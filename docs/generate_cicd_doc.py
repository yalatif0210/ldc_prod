from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page setup ────────────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Cm(21)
section.page_height = Cm(29.7)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)

# ── Palette ───────────────────────────────────────────────────────────────────
COLOR_TITLE   = RGBColor(0x17, 0x88, 0x8F)   # teal brand
COLOR_H1      = RGBColor(0x1F, 0x29, 0x37)   # dark
COLOR_H2      = RGBColor(0x17, 0x88, 0x8F)
COLOR_H3      = RGBColor(0x2E, 0x86, 0xAB)
COLOR_CODE_BG = RGBColor(0xF3, 0xF4, 0xF6)
COLOR_CODE_FG = RGBColor(0x1F, 0x29, 0x37)
COLOR_TH_BG   = RGBColor(0x17, 0x88, 0x8F)
COLOR_TH_FG   = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TR_ALT  = RGBColor(0xF0, 0xFD, 0xFD)

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(22)
    run.font.color.rgb = COLOR_TITLE
    p.space_after = Pt(4)

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    run.italic = True
    p.space_after = Pt(16)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(15)
    run.font.color.rgb = COLOR_H1
    # bottom border
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '17888F')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_H2

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_H3

def add_body(doc, text, bold_parts=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text).font.size = Pt(10.5)

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(10.5)

def add_code(doc, lines):
    """Bloc de code avec fond gris."""
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(1)
        run = p.add_run(line if line else " ")
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_CODE_FG
        # fond gris
        pPr  = p._p.get_or_add_pPr()
        shd  = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  'F3F4F6')
        pPr.append(shd)

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # header
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TH_FG
        set_cell_bg(cell, COLOR_TH_BG)
    # rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        bg  = COLOR_TR_ALT if r_idx % 2 == 1 else RGBColor(0xFF,0xFF,0xFF)
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
            set_cell_bg(cell, bg)
    doc.add_paragraph()   # espace après le tableau

def add_note(doc, text, color=RGBColor(0x17,0x88,0x8F)):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(8)
    run = p.add_run('ℹ  ' + text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = color

def page_break(doc):
    doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# CONTENU DU DOCUMENT
# ══════════════════════════════════════════════════════════════════════════════

add_title(doc, "Guide CI/CD — LDC Application")
add_subtitle(doc, "Configuration du pipeline, déploiement et partage des images Docker")

# ── 1. Introduction ───────────────────────────────────────────────────────────
add_h1(doc, "1. Introduction")
add_body(doc,
    "Ce document décrit la mise en place complète d'un pipeline CI/CD pour une application "
    "full-stack composée d'un frontend Angular, d'un backend Spring Boot et de deux bots Python. "
    "Le pipeline utilise GitHub Actions comme moteur d'automatisation et GitHub Container Registry "
    "(ghcr.io) comme registre d'images Docker."
)

add_h2(doc, "1.1  Services de l'application")
add_table(doc,
    ["Service", "Technologie", "Répertoire source"],
    [
        ["Frontend",    "Angular",              "./ldc_frontend"],
        ["Backend",     "Spring Boot (Java 17)", "./ldc_backend"],
        ["Support Bot", "Flask (Python)",        "./support_ldc"],
        ["Pilot Bot",   "Python",               "./pilot_bot"],
    ]
)

add_h2(doc, "1.2  Prérequis")
add_bullet(doc, "Un compte GitHub avec accès au dépôt")
add_bullet(doc, "Un VPS Linux avec Docker et Docker Compose installés")
add_bullet(doc, "Un nom de domaine (pour la production)")
add_bullet(doc, "Un certificat SSL (Let's Encrypt ou autre)")

# ── 2. Architecture ───────────────────────────────────────────────────────────
page_break(doc)
add_h1(doc, "2. Architecture générale")
add_body(doc,
    "Le pipeline repose sur deux branches principales et deux environnements distincts :"
)
add_code(doc, [
    "Code local",
    "    │",
    "    ├── develop ──► push ──► CI (lint + build) ──► CD Staging (automatique)",
    "    │",
    "    └── main ────► tag vX.Y.Z ──► CI (lint + build) ──► CD Production (approbation manuelle)",
])

add_h2(doc, "2.1  Environnements")
add_table(doc,
    ["Environnement", "Branche", "Tag image", "Port", "Déclencheur"],
    [
        ["Staging",    "develop", "staging-XXXXXXX", "8081", "Automatique au push"],
        ["Production", "main (tags)", "vX.Y.Z",      "443",  "Approbation manuelle"],
    ]
)

add_h2(doc, "2.2  Fichiers Docker Compose")
add_table(doc,
    ["Fichier", "Rôle"],
    [
        ["docker-compose.yml",         "Base commune à tous les environnements"],
        ["docker-compose.staging.yml", "Surcharges staging : ports, noms de conteneurs"],
        ["docker-compose.prod.yml",    "Surcharges production : port 443"],
        ["docker-compose.ci.yml",      "Utilise les images pré-buildées depuis ghcr.io"],
    ]
)

# ── 3. Mise en place GitHub Actions ───────────────────────────────────────────
page_break(doc)
add_h1(doc, "3. Mise en place du pipeline GitHub Actions")

add_h2(doc, "3.1  Configurer les secrets du dépôt")
add_body(doc,
    "Dans GitHub, aller dans : Settings → Secrets and variables → Actions → New repository secret"
)
add_table(doc,
    ["Nom du secret", "Description", "Exemple"],
    [
        ["VPS_HOST",         "Adresse IP du VPS",                  "207.180.209.55"],
        ["VPS_USER",         "Utilisateur SSH",                    "root"],
        ["VPS_SSH_KEY",      "Clé SSH privée (contenu complet)",   "-----BEGIN OPENSSH..."],
        ["VPS_PORT",         "Port SSH",                           "22"],
        ["VPS_PROD_PATH",    "Chemin de la prod sur le VPS",       "/opt/ldc_prod"],
        ["VPS_STAGING_PATH", "Chemin du staging sur le VPS",       "/opt/ldc_staging"],
    ]
)

add_h2(doc, "3.2  Configurer l'environnement de production")
add_body(doc,
    "L'approbation manuelle avant déploiement en production est gérée par les "
    "GitHub Environments."
)
add_bullet(doc, "Aller dans : Settings → Environments → New environment")
add_bullet(doc, "Nommer l'environnement : production")
add_bullet(doc, "Cocher : Required reviewers")
add_bullet(doc, "Ajouter son propre compte comme approbateur")
add_note(doc,
    "Sans cet environnement configuré, le déploiement en production se ferait "
    "automatiquement sans validation humaine."
)

add_h2(doc, "3.3  Structure des workflows")
add_body(doc, "Les fichiers de workflow sont dans le répertoire .github/workflows/ :")
add_table(doc,
    ["Fichier", "Déclencheur", "Rôle"],
    [
        ["ci.yml",            "Tout push",         "Lint + Build (validation du code)"],
        ["cd-staging.yml",    "Push sur develop",  "Build images → Deploy staging"],
        ["cd-production.yml", "Tag vX.Y.Z sur main","Build images → Deploy production"],
    ]
)

# ── 4. Pipeline CI ─────────────────────────────────────────────────────────────
add_h1(doc, "4. Pipeline CI — Intégration continue")
add_body(doc,
    "Le CI s'exécute sur chaque push, toutes branches confondues. "
    "Il valide que le code compile et respecte les règles de qualité. "
    "Si le CI échoue, le CD ne se déclenche pas."
)

add_h2(doc, "4.1  Jobs du CI")
add_table(doc,
    ["Job", "Outil", "Ce qui est vérifié"],
    [
        ["lint-frontend",  "ng lint + stylelint",        "TypeScript, SCSS"],
        ["build-frontend", "ng build --configuration=production", "Compilation Angular AOT"],
        ["build-backend",  "mvn package -DskipTests",    "Compilation Java"],
        ["lint-python",    "flake8",                     "Style Python (support_ldc + pilot_bot)"],
    ]
)

add_h2(doc, "4.2  Points d'attention")
add_bullet(doc, "persist-credentials: false sur tous les checkouts pour éviter l'erreur git exit 128")
add_bullet(doc, "FORCE_JAVASCRIPT_ACTIONS_TO_NODE24: true pour éviter les warnings de dépréciation Node.js 20")
add_bullet(doc, "Les fichiers vendor (OpenLMIS) sont exclus du lint SCSS via ignoreFiles dans .stylelintrc")

# ── 5. Pipeline CD Staging ────────────────────────────────────────────────────
page_break(doc)
add_h1(doc, "5. Pipeline CD — Staging (automatique)")
add_body(doc,
    "Tout push sur develop déclenche automatiquement un build et un déploiement sur le serveur staging."
)

add_h2(doc, "5.1  Étapes du pipeline staging")
add_bullet(doc, "Étape 1 — Checkout du code source")
add_bullet(doc, "Étape 2 — Connexion à GitHub Container Registry (ghcr.io)")
add_bullet(doc, "Étape 3 — Build et push des 4 images avec les tags :")
add_bullet(doc, "staging  (tag fixe, pointe toujours sur le dernier build)", level=1)
add_bullet(doc, "staging-XXXXXXX  (SHA court du commit, pour traçabilité)", level=1)
add_bullet(doc, "Étape 4 — Déploiement SSH sur le VPS :")
add_code(doc, [
    "cd /opt/ldc_staging",
    "git fetch origin develop && git checkout origin/develop",
    "docker login ghcr.io",
    "docker compose -f docker-compose.yml \\",
    "               -f docker-compose.staging.yml \\",
    "               -f docker-compose.ci.yml pull",
    "docker compose -f docker-compose.yml \\",
    "               -f docker-compose.staging.yml \\",
    "               -f docker-compose.ci.yml up -d --no-build --remove-orphans",
    "docker image prune -f",
])
add_note(doc,
    "docker compose down n'est jamais exécuté en staging afin de préserver "
    "les données de la base de données (volumes Docker)."
)

# ── 6. Pipeline CD Production ─────────────────────────────────────────────────
add_h1(doc, "6. Pipeline CD — Production (approbation manuelle)")
add_body(doc,
    "La mise en production est déclenchée par la création d'un tag Git au format vX.Y.Z. "
    "Le déploiement ne s'effectue qu'après approbation manuelle dans GitHub Actions."
)

add_h2(doc, "6.1  Étapes du pipeline production")
add_bullet(doc, "Étape 1 — Build et push des 4 images avec les tags :")
add_bullet(doc, "latest", level=1)
add_bullet(doc, "vX.Y.Z  (ex: v1.0.4)", level=1)
add_bullet(doc, "Étape 2 — Attente d'approbation humaine (GitHub Environment production)")
add_bullet(doc, "Étape 3 — Déploiement SSH sur le VPS :")
add_code(doc, [
    "cd /opt/ldc_prod",
    "git fetch --tags origin && git checkout vX.Y.Z",
    "docker login ghcr.io",
    "docker compose -f docker-compose.yml \\",
    "               -f docker-compose.prod.yml \\",
    "               -f docker-compose.ci.yml pull",
    "docker compose -f docker-compose.yml \\",
    "               -f docker-compose.prod.yml \\",
    "               -f docker-compose.ci.yml up -d --no-build",
    "docker image prune -f",
])

add_h2(doc, "6.2  Comment approuver un déploiement")
add_bullet(doc, "Aller sur GitHub → dépôt → onglet Actions")
add_bullet(doc, "Cliquer sur le run CD — Production en attente")
add_bullet(doc, "Cliquer sur Review deployments")
add_bullet(doc, "Sélectionner l'environnement production → Approve and deploy")

# ── 7. Procédure develop ↔ main ───────────────────────────────────────────────
page_break(doc)
add_h1(doc, "7. Procédure de basculement develop ↔ main")

add_h2(doc, "7.1  Travail quotidien (mode développement)")
add_body(doc,
    "On travaille toujours sur la branche develop. Pour pousser des modifications "
    "vers staging, utiliser le script utilitaire :"
)
add_code(doc, ["bash scripts/push-develop.sh \"description de la modification\""])
add_body(doc, "Le staging se déploie automatiquement après le push.")

add_h2(doc, "7.2  Mise en production (release)")

add_h3(doc, "Étape 1 — S'assurer que develop est propre")
add_code(doc, [
    "git checkout develop",
    "git pull origin develop",
])

add_h3(doc, "Étape 2 — Passer sur main et merger develop")
add_code(doc, [
    "git checkout main",
    "git pull origin main",
    "git merge develop --no-edit",
])

add_h3(doc, "Étape 3 — Releaser la version")
add_code(doc, [
    "git stash                          # si des fichiers locaux non commités",
    "bash scripts/bump-version.sh X.Y.Z # ex: bash scripts/bump-version.sh 1.1.0",
    "git stash pop",
])
add_body(doc, "Le script bump-version.sh effectue automatiquement :")
add_bullet(doc, "Mise à jour de package.json, package-lock.json et pom.xml")
add_bullet(doc, "Création d'un commit chore: release vX.Y.Z")
add_bullet(doc, "Création du tag vX.Y.Z")
add_bullet(doc, "Push vers origin/main avec le tag")

add_h3(doc, "Étape 4 — Approuver le déploiement")
add_body(doc, "Voir section 6.2 — Approuver un déploiement.")

add_h3(doc, "Étape 5 — Revenir en mode développement")
add_code(doc, [
    "git checkout develop",
    "git merge main --no-edit   # ramener le bump de version dans develop",
    "git push origin develop",
])

add_h2(doc, "7.3  Schéma du cycle complet")
add_code(doc, [
    "develop  ──●──●──●──────────────────────────●──►",
    "                    ↘ merge                  ↑",
    "main     ────────────●── tag vX.Y.Z ─────────●──►",
    "                                         merge back",
])

add_h2(doc, "7.4  Règles importantes")
add_table(doc,
    ["Règle", "Raison"],
    [
        ["Ne jamais commiter directement sur main",       "Seul bump-version.sh touche main"],
        ["Toujours merger main → develop après release",  "Évite les conflits de version"],
        ["Tester en staging avant de releaser",           "La prod ne reçoit que du code validé"],
        ["git stash avant bump-version.sh si tree sale",  "Le script exige un working tree propre"],
    ]
)

add_h2(doc, "7.5  Numérotation des versions (SemVer)")
add_table(doc,
    ["Type de changement", "Exemple"],
    [
        ["Correctif (bug fix)",       "1.0.3 → 1.0.4"],
        ["Nouvelle fonctionnalité",   "1.0.4 → 1.1.0"],
        ["Changement majeur (breaking change)", "1.1.0 → 2.0.0"],
    ]
)

# ── 8. Retrouver une ancienne version ─────────────────────────────────────────
page_break(doc)
add_h1(doc, "8. Retrouver et restaurer une ancienne version")

add_h2(doc, "8.1  Consulter le code d'une ancienne version")
add_code(doc, [
    "# Voir un fichier spécifique au tag v1.0.2",
    "git show v1.0.2:ldc_frontend/src/app/routes/public/report/lab-report/lab-report.ts",
    "",
    "# Naviguer dans tout le code du tag v1.0.2",
    "git checkout v1.0.2",
    "# ... consulter les fichiers ...",
    "git checkout develop   # revenir",
])

add_h2(doc, "8.2  Redéployer une ancienne version en production")
add_body(doc,
    "La méthode recommandée est de créer un nouveau tag pointant sur l'ancien commit. "
    "On n'annule jamais l'historique, on avance."
)
add_code(doc, [
    "# Redéployer la v1.0.2 en créant le tag v1.0.5 sur le même commit",
    "git tag v1.0.5 v1.0.2",
    "git push origin v1.0.5",
    "# GitHub Actions déclenche le build → approuver → prod revient à l'état v1.0.2",
])

add_h2(doc, "8.3  Récupérer un fichier d'une ancienne version")
add_code(doc, [
    "git checkout v1.0.2 -- ldc_frontend/src/app/routes/public/report/lab-report/lab-report.ts",
    "# Le fichier revient à l'état v1.0.2 dans le working tree, prêt à être commité",
])

add_h2(doc, "8.4  Comparer deux versions")
add_code(doc, [
    "git diff v1.0.2 v1.0.4                          # tous les changements",
    "git log --oneline v1.0.2..v1.0.4                # liste des commits entre les deux",
    "git tag --sort=-version:refname                  # toutes les versions (plus récent en premier)",
])

# ── 9. Partage des images pour déploiement ────────────────────────────────────
page_break(doc)
add_h1(doc, "9. Partager les images pour déploiement chez un tiers")
add_body(doc,
    "Il est possible de fournir l'application prête à l'emploi à un tiers "
    "sans lui donner accès au code source. Il n'a besoin que de Docker."
)

add_h2(doc, "9.1  Ce que le tiers reçoit")
add_table(doc,
    ["Élément", "Partagé ?", "Remarque"],
    [
        ["Code source",            "Non",  "Reste privé dans le dépôt Git"],
        ["Images Docker (ghcr.io)","Oui",  "Via registry public ou token d'accès"],
        ["docker-compose.yml",     "Oui",  "Aucun code source, uniquement la config"],
        ["docker-compose.prod.yml","Oui",  "Config port 443"],
        ["docker-compose.ci.yml",  "Oui",  "Référence les images du registry"],
        [".env",                   "Oui",  "Avec ses propres valeurs (pas les vôtres)"],
    ]
)

add_h2(doc, "9.2  Option A — Images publiques (le plus simple)")
add_body(doc,
    "Rendre les images accessibles sans authentification :"
)
add_bullet(doc, "GitHub → onglet Packages → cliquer sur chaque image")
add_bullet(doc, "Package settings → Change visibility → Public")
add_body(doc, "Le tiers n'a besoin d'aucun compte GitHub. Il exécute simplement :")
add_code(doc, [
    "export IMAGE_TAG=v1.0.4",
    "export CI_IMAGE_OWNER=yalatif0210",
    "docker compose -f docker-compose.yml -f docker-compose.prod.yml -f docker-compose.ci.yml pull",
    "docker compose -f docker-compose.yml -f docker-compose.prod.yml -f docker-compose.ci.yml up -d --no-build",
])

add_h2(doc, "9.3  Option B — Images privées (token d'accès)")
add_body(doc,
    "Conserver les images privées et fournir un token GitHub avec la permission read:packages :"
)
add_bullet(doc, "GitHub → Settings → Developer settings → Personal access tokens")
add_bullet(doc, "Générer un token avec la permission read:packages uniquement")
add_bullet(doc, "Transmettre ce token au tiers")
add_body(doc, "Le tiers se connecte au registry avec ce token :")
add_code(doc, [
    "echo \"LE_TOKEN\" | docker login ghcr.io -u yalatif0210 --password-stdin",
    "",
    "export IMAGE_TAG=v1.0.4",
    "export CI_IMAGE_OWNER=yalatif0210",
    "docker compose -f docker-compose.yml -f docker-compose.prod.yml -f docker-compose.ci.yml pull",
    "docker compose -f docker-compose.yml -f docker-compose.prod.yml -f docker-compose.ci.yml up -d --no-build",
])

add_h2(doc, "9.4  Le fichier .env du tiers")
add_body(doc,
    "Le tiers doit créer son propre fichier .env avec ses valeurs. "
    "Fournir le template suivant sans les valeurs sensibles :"
)
add_code(doc, [
    "# .env — à remplir avec vos propres valeurs",
    "DB_PASSWORD=",
    "JWT_SECRET=",
    "APP_CORS_ALLOWED_ORIGINS=https://son-domaine.com",
    "TELEGRAM_BOT_TOKEN=",
    "TELEGRAM_BOT_USERNAME=",
])
add_note(doc,
    "Ne jamais transmettre votre propre fichier .env. "
    "Chaque déploiement doit avoir ses propres secrets."
)

add_h2(doc, "9.5  Prérequis sur le VPS du tiers")
add_bullet(doc, "Docker Engine installé (version 24+)")
add_bullet(doc, "Docker Compose plugin installé (version 2+)")
add_bullet(doc, "Ports 80 et 443 ouverts dans le pare-feu")
add_bullet(doc, "Un certificat SSL configuré (Let's Encrypt recommandé)")

# ── 10. Commandes de maintenance ──────────────────────────────────────────────
page_break(doc)
add_h1(doc, "10. Commandes de maintenance sur le VPS")

add_h2(doc, "10.1  Vérifier l'état des conteneurs")
add_code(doc, [
    "docker ps --format \"table {{.Names}}\\t{{.Image}}\\t{{.Status}}\"",
])

add_h2(doc, "10.2  Consulter les logs")
add_code(doc, [
    "docker logs ldc_backend --tail=100 -f",
    "docker logs staging_ldc_backend --tail=100 -f",
    "docker logs ldc_frontend --tail=50",
])

add_h2(doc, "10.3  Redémarrer un service")
add_code(doc, [
    "docker restart ldc_frontend",
    "docker restart ldc_backend",
])

add_h2(doc, "10.4  Vérifier les variables d'environnement")
add_code(doc, [
    "docker exec ldc_backend env | grep CORS",
    "docker exec staging_ldc_backend env | grep CORS",
])

add_h2(doc, "10.5  Libérer l'espace disque")
add_code(doc, [
    "docker image prune -f          # supprimer les images inutilisées",
    "docker system prune -f         # supprimer tout ce qui n'est pas utilisé",
])

# ── Sauvegarde ────────────────────────────────────────────────────────────────
output_path = r"c:\DEV_APP_LAB\DEPLOY\app_v3_claude_production\docs\CICD_Guide.docx"
doc.save(output_path)
print(f"Document généré : {output_path}")
